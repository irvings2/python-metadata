from openpyxl import load_workbook
from  PyPDF2 import *
import docx
import sys
import os

def imprimirMetadatos(target):
    if not os.path.isdir(target):
        print(f"No existe el directorio {target}")
        return
    walk = os.walk(target)
    for rutadir, nombresdir, archivos in walk:
        for nombre in archivos:
            extension = nombre.lower().rsplit(".",1)[-1]
            rutaCompletaArchivo = rutadir + os.path.sep + nombre
            if extension == "xlsx":
                imprimirXlsx(rutaCompletaArchivo)
            elif extension == "pdf":
                imprimirPdf(rutaCompletaArchivo)
            elif extension == "docx":
                imprimirDocx(rutaCompletaArchivo)
            else:
                pass

def imprimirXlsx(rutaCompletaArchivo):
    print(f"Metadatos del archivo {rutaCompletaArchivo}")
    xlsx_file = load_workbook(rutaCompletaArchivo)
    print(f"Title: {xlsx_file.properties.title}")
    print(f"Creator: {xlsx_file.properties.creator}")
    print(f"Description: {xlsx_file.properties.description}")
    print(f"Subject: {xlsx_file.properties.subject}")
    print(f"Identifier: {xlsx_file.properties.identifier}")
    print(f"Language: {xlsx_file.properties.language}")
    print(f"Created: {xlsx_file.properties.created}")
    print(f"Modified: {xlsx_file.properties.modified}")
    print(f"Last Modified By: {xlsx_file.properties.lastModifiedBy}")
    print(f"Revision: {xlsx_file.properties.revision}")
    print(f"Keywords: {xlsx_file.properties.keywords}")
    print(f"Category: {xlsx_file.properties.category}")
    print(f"Content Status: {xlsx_file.properties.contentStatus}")
    print(f"Last Printed: {xlsx_file.properties.lastPrinted}")
    print("")

def imprimirPdf(rutaCompletaArchivo):
    print(f"Metadatos del archivo {rutaCompletaArchivo}")
    pdf_file = PdfReader(rutaCompletaArchivo)
    pdf_info = pdf_file.metadata
    for metaItem in pdf_info:
        print(metaItem[1:] + ": " + pdf_info[metaItem])
    print("")

def imprimirDocx(rutaCompletaArchivo):
    print(f"Metadatos del archivo {rutaCompletaArchivo}")
    docx_file = docx.Document(rutaCompletaArchivo)
    prop = docx_file.core_properties

    attrs = ["author", "category", "comments", "content_status", 
         "created", "identifier", "keywords", "language", 
         "last_modified_by", "last_printed", "modified", 
         "revision", "subject", "title", "version"]

    for attr in attrs:
        value = getattr(prop, attr)
        if value:
            print(f"{attr}: {value}")
    print("")

def main(argv):
    if len(argv) != 2:
        print("Ingrese los argumentos necesarios")
        return
    else:
        target = argv[1]
        imprimirMetadatos(target)

if __name__ == "__main__":
    main(sys.argv)

